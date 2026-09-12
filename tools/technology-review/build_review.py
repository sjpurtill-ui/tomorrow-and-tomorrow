from pathlib import Path
import json, re, sys, collections
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont
from design_content import PAGES,p,h,bullets,table
from branches_content import EXTRA_PAGES, MILITARY_ROLES
from atlas_content import DOMAINS

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/technology-review'
QA=ROOT/'artifacts/technology-review'
OUT.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)

# This is an original explanatory diagram, not an edit of the supplied art.
im=Image.new('RGB',(1700,790),'white');draw=ImageDraw.Draw(im)
font_path='/System/Library/Fonts/Supplemental/Arial.ttf'
font=ImageFont.truetype(font_path,27); small=ImageFont.truetype(font_path,23)
boxes={
 'a':(20,40,410,160,'Airflow and spoilage\nobservations'),
 'b':(20,230,410,350,'Controlled fire and\nfood observations'),
 'c':(20,440,410,570,'Known foreign practice\nand visiting scholar'),
 'd':(530,65,950,175,'Food drying'),
 'e':(530,255,950,365,'Smoke preservation'),
 'f':(1180,145,1680,320,'Preserved provisions\nusing a suitable method'),
 'g':(550,580,990,700,'Containers and surplus\nplus organized carrying'),
 'j':(1190,550,1670,710,'Longer viable journeys\nwith actual supplies')}
for key,(x1,y1,x2,y2,txt) in boxes.items():
    draw.rounded_rectangle((x1,y1,x2,y2),radius=10,fill='#F2F1EC',outline='#686860',width=2)
    draw.multiline_text(((x1+x2)/2,(y1+y2)/2),txt,font=font,fill='#191919',anchor='mm',align='center',spacing=8)
def arrow(points,label=None,where=None,dashed=False):
    if dashed:
        for a,b in zip(points,points[1:]):
            n=max(1,int(((b[0]-a[0])**2+(b[1]-a[1])**2)**.5/12))
            for i in range(0,n,2):
                c=(a[0]+(b[0]-a[0])*i/n,a[1]+(b[1]-a[1])*i/n)
                d=(a[0]+(b[0]-a[0])*min(i+1,n)/n,a[1]+(b[1]-a[1])*min(i+1,n)/n)
                draw.line([c,d],fill='#4E6566',width=3)
    else:draw.line(points,fill='#4E6566',width=3)
    x,y=points[-1];ax,ay=points[-2]
    if abs(x-ax)>abs(y-ay):draw.polygon([(x,y),(x-12,y-7),(x-12,y+7)],fill='#4E6566')
    else:draw.polygon([(x,y),(x-7,y-12),(x+7,y-12)],fill='#4E6566')
    if label:draw.text(where,label,font=small,fill='#191919')
arrow([(410,100),(530,120)])
arrow([(410,290),(530,310)])
arrow([(410,495),(480,495),(480,345),(530,345)],'Teaching',(490,438),True)
arrow([(950,120),(1080,120),(1080,200),(1180,200)])
arrow([(950,310),(1080,310),(1080,260),(1180,260)])
draw.text((1093,222),'OR',font=small,fill='#191919')
arrow([(1430,320),(1430,550)],'AND',(1452,440))
arrow([(990,640),(1190,640)])
draw.text((20,738),'Solid lines: causal capability links     Dashed line: acquired evidence and teaching',font=small,fill='#444444')
im.save(QA/'branch-diagram.png')

pages=list(PAGES)
def insert_before(title,items):
    idx=next(i for i,x in enumerate(pages) if x[0]==title)
    pages[idx:idx]=items

insert_before('Making dependence measurable',[EXTRA_PAGES[0]])
insert_before('From an observation to a changed society',[EXTRA_PAGES[1]])
insert_before('The simulation architecture that supports the breadth',EXTRA_PAGES[2:])

scope_page=('Coverage across the whole historical span',[
 p('The 5,000 target is a single set of distinct discoveries, counted once by primary editorial field and once in a separate historical coverage view. Those two views describe the same catalog; they are not multiplied together. Geography, teaching routes, variants, and proficiency states do not increase the count.'),
 table(['Editorial horizon','Target','Examples of required depth'],[
 ['Survival and early practice','600','Food, fire, stone, fibers, care, movement, memory, seasonal knowledge'],
 ['Settlement and regional exchange','700','Water, crops, animals, building, metals, exchange, local institutions'],
 ['Complex societies and learned traditions','800','Cities, sea routes, mathematics, records, law, medicine, organized crafts'],
 ['Mechanization and industrial systems','700','Power, precision, chemicals, rail, sanitation, mass production and coordination'],
 ['Modern scientific and networked systems','900','Electricity, aviation, computation, medicine, communications, planetary measurement'],
 ['Advanced planetary and human futures','1,300','Frontier energy, biology, automation, resilient institutions, oceans, space industry and habitation'],
 ['Total distinct discoveries','5,000','Authoring allocation for review, not completed content']],[2.4,.65,3.75]),
 p('These are proposed coverage allocations. Each category needs internal diversity and enough causal depth to feel like generations of work. The future allocation includes engineering and institutional development grounded in known principles, not 1,300 miraculous inventions.'),
 h('The atlas is a starting set'),
 p('The following field allocations reserve the full scope. Later in this document, the 600 candidate subjects demonstrate its breadth. They are deliberately uneven in historical starting point: there is no prehistoric semiconductor program. Reading groups do not impose dates or dependency edges.'),
 p('The complete authoring catalog must expand these subjects into genuinely distinct capabilities with reviewed relationships and consequences. The present document does not pretend that a list of names alone is an implemented technology system.')])

coverage_pages=[scope_page]
for start in (0,12):
    rows=[[d['id'],d['name'],str(d['target'])] for d in DOMAINS[start:start+12]]
    coverage_pages.append(('Fields of discovery '+('from food to measurement' if start==0 else 'from communication to human futures'),[
        p('These 24 editorial fields organize authoring and review. They need not replace the existing twelve simulation domains or become 24 mandatory player allocation sliders. Cross-field capabilities have one primary identity and explicit connections.'),
        table(['Field','Primary subject','Target discoveries'],rows,[.6,4.7,1.5]),
        p('Target allocation across all 24 fields: 5,000 distinct discoveries. The review atlas contributes 25 named candidates in each field, 600 total. Counts are scope commitments to evaluate; they are not a reason to keep duplicate or mechanically empty entries.'),
        p('A subject can support several fields without being counted several times. Materials knowledge supports armor, bridges, power plants, and spacecraft; the applications can be separate discoveries when they introduce a distinct capability or mechanism.')]))
insert_before('How the technology web should look and feel',coverage_pages)

# Keep review decisions before the appendix, with each atlas page carrying two fields.
for start in range(0,24,2):
    blocks=[p('Candidate subjects for review. Groups aid reading; they are not prerequisite chains, fixed eras, or claims of historical first invention. Frontier entries require individual confidence and engineering review.')]
    for d in DOMAINS[start:start+2]:
        blocks += [h(d['id']+' '+d['name']),p(d['purpose']+' Proposed full field allocation: '+str(d['target'])+' discoveries.')]
        for i,group in enumerate(d['groups']):
            blocks.append(('atlas',f"{d['id']} {i*5+1:02d} to {i*5+5:02d}",group+'.'))
        blocks.append(('small','Connections: '+d['links']))
    pages.append(('Candidate atlas '+DOMAINS[start]['id']+' and '+DOMAINS[start+1]['id'],blocks))

sources=[
 ('R1','Current discovery system','scripts/discovery_system.gd at 940d5a2. Initialization selects non-frontier entries for the live catalog; _path_is_viable excludes frontier entries.',None),
 ('R2','Legacy frontier catalog','scripts/discovery_frontier_catalog.gd at 940d5a2. Defines 4,608 combinations and maturity stages through day 1,095,000.',None),
 ('R3','Live inventory','197 entries exported from the initialized technology_catalog in the isolated artwork worktree based on 940d5a2. Latest earliest gate: day 98,000. Security entries: 58.',None),
 ('R4','Retirement checkpoint','Git checkpoint bac20351564abc387f03c8dbd42dd0ed97a6055d, September 4 2026; README section Concrete technology branches.',None),
 ('R5','Branching and knowledge exchange','scripts/knowledge_pathways.gd and docs/KNOWLEDGE_EXCHANGE_HANDOFF.md. Fifteen alternative routes, one discovery/adoption record, actual evidence and travel.',None),
 ('R6','Long campaign and world constraints','docs/WHOLE_GAME_COHERENCE.md and docs/GAMEPLAY_COHERENCE_ROADMAP.md. Long horizon, no mandatory era ladder, thousands of routes, shared rules, bounded scale.',None),
 ('R7','Military authority','AGENTS.md and docs/GENERAL_CAMPAIGN_DESIGN.md. Generals execute operations; the player directs objectives through conversation.',None),
 ('R8','User direction','This conversation, September 10 2026: full-history scope; disadvantaged recovery routes; paid or traded scholar envoys; later purchased research; broader military units and research; approved visual references.',None),
]
pages.append(('Repository evidence and design provenance',[
 p('Repository findings describe the current source at the stated revision. Design recommendations elsewhere in the document are proposals. No game change, graph migration, or new military capability is claimed implemented by this document.'),
 *[p(f'[{code}] {title}. {desc}') for code,title,desc,url in sources],
 p('The candidate atlas is manually authored design content. Its 600 IDs are review identifiers, not replacement save IDs. Existing production IDs require explicit reconciliation before any implementation.')]))

external=[
 ('H1','Smithsonian Human Origins Program','Stone Tools','https://humanorigins.si.edu/evidence/behavior/stone-tools','Supports the distinction between deep human technological history and a compressed 3,000-year campaign.'),
 ('H2','UNESCO','200 years before Gutenberg The master printers of Koryo','https://www.unesco.org/en/articles/200-years-gutenberg-master-printers-koryo-0','Supports including multiple printing traditions rather than a single European sequence.'),
 ('F1','NASA','Technology Readiness Levels','https://www.nasa.gov/directorates/somd/space-communications-navigation-program/technology-readiness-levels/','Supports separating a concept, demonstration, and mature operating technology.'),
 ('F2','ITER','Making fusion work','https://www.iter.org/fusion-energy/making-it-work','Identifies integration, material, fuel-cycle, and heat-removal challenges for useful fusion systems.'),
 ('F2a','ITER','Frequently asked questions','https://www.iter.org/faqs','Confirms that ITER itself is not designed to generate electricity.'),
 ('F3','NIST','Quantum Computing Explained','https://www.nist.gov/quantum-information-science/quantum-computing-explained','Supports workload-specific capabilities and the importance of reliable operation rather than a universal computing bonus.'),
 ('F4','NASA','Mars Oxygen In Situ Resource Utilization Experiment','https://www.nasa.gov/space-technology-mission-directorate/tdm/mars-oxygen-in-situ-resource-utilization-experiment-moxie/','Supports treating in situ oxygen production as a demonstrated component, distinct from a complete settlement.'),
 ('F5','NASA','2024 NASA Technology Taxonomy','https://www.nasa.gov/otps/2024-nasa-technology-taxonomy/','Additional authoring reference for the breadth of space technology disciplines and supporting systems.'),
]
pages.append(('Historical and engineering references',[
 p('Selected primary institutional sources consulted for this review. These ground the distinctions cited in the text, not every candidate entry. The full catalog requires node-level sourcing, uncertainty labels, and review of disputed historical dependencies.'),
 *[('source',code,org,title,url,note) for code,org,title,url,note in external],
 p('All future timing, cost examples, field allocations, role inventories, and balance targets in this proposal are game design assumptions rather than predictions or measured game results.')]))

contents_titles=[
 'How to review this proposal','The verified gap in the current game','The branching contract',
 'The routes to the same technology','Scholar envoys and purchased research',
 'The structure of a discovery','Campaign time and historical breadth',
 'A substantial future with honest uncertainty','Coverage across the whole historical span',
 'How the technology web should look and feel','Worked branches for food and water',
 'Military forces built from capabilities','Military research is more than weapons',
 'The simulation architecture that supports the breadth','Acceptance criteria for the complete system',
 'Decisions for this review','Candidate atlas D01 and D02','Repository evidence and design provenance']
pages.insert(1,('Contents',[
 p('Read the branching and acquisition sections first. The worked branches show the proposed rules in practice. The military design and candidate atlas provide the broader content inventory.'),
 ('contents',contents_titles),
 p('All section titles also appear in Word’s navigation pane. The catalog allocation is a proposed full-production scope; the atlas entries and military roles are concrete review candidates.')]))

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.10
normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',26),('Subtitle',12),('Heading 1',20),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.font.bold=name!='Subtitle'
    st.paragraph_format.space_before=Pt(9 if name=='Heading 2' else 0)
    st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.keep_with_next=True
for st in ['List Bullet','List Number']:
    doc.styles[st].font.name='Arial';doc.styles[st].font.size=Pt(11)
    doc.styles[st].paragraph_format.space_after=Pt(7)
    doc.styles[st].paragraph_format.line_spacing=1.08
header=sec.header.paragraphs[0]
header.text='TOMORROW AND TOMORROW   /   TECHNOLOGY DESIGN REVIEW'
header.runs[0].font.name='Arial';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Review draft 1   •   ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Technology and civilization across all history'
doc.core_properties.subject='Review proposal for the technology and military research systems of Tomorrow and Tomorrow'
doc.core_properties.author='Codex'

def bookmark(par,name,num):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(num));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(num))
    par._p.insert(0,a);par._p.append(b)

def link(par,label,url=None,anchor=None):
    element=OxmlElement('w:hyperlink')
    if url:element.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
    if anchor:element.set(qn('w:anchor'),anchor)
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'274C59');pr.append(col)
    uu=OxmlElement('w:u');uu.set(qn('w:val'),'single');pr.append(uu);rr.append(pr)
    tt=OxmlElement('w:t');tt.text=label;rr.append(tt);element.append(rr);par._p.append(element)

def add_table(headers,rows,widths):
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    if not widths:widths=[6.8/len(headers)]*len(headers)
    for j,w in enumerate(widths):tb.columns[j].width=Inches(w)
    for row_i,vals in enumerate([headers]+rows):
        row=tb.rows[0] if row_i==0 else tb.add_row()
        prop=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');prop.append(no_split)
        if row_i==0:
            repeat=OxmlElement('w:tblHeader');prop.append(repeat)
        for j,value in enumerate(vals):
            cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'364A50' if row_i==0 else ('F2F4F4' if row_i%2==0 else 'FFFFFF'));tcpr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');mar.append(el)
            tcpr.append(mar)
            pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.04
            rr=pp.add_run(str(value));rr.font.name='Arial';rr.font.size=Pt(10)
            if row_i==0:rr.bold=True;rr.font.color.rgb=RGBColor(255,255,255)
            if str(value).replace(',','').isdigit():pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(1);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(2)

markdown=[]
for i,(title,blocks) in enumerate(pages):
    if i:doc.add_page_break()
    hp=doc.add_paragraph(title,'Title' if i==0 else 'Heading 1');bookmark(hp,'section_'+str(i),i+1)
    markdown.append('# '+title+'\n')
    for b in blocks:
        kind=b[0]
        if kind=='p':doc.add_paragraph(b[1]);markdown.append(b[1]+'\n')
        elif kind=='h':doc.add_paragraph(b[1],'Heading 2');markdown.append('## '+b[1]+'\n')
        elif kind=='bullets':
            for text in b[1]:doc.add_paragraph(text,'List Bullet');markdown.append('- '+text)
            markdown.append('')
        elif kind=='table':
            add_table(b[1],b[2],b[3]);markdown.append('| '+' | '.join(b[1])+' |');markdown.append('| '+' | '.join(['---']*len(b[1]))+' |')
            for row in b[2]:markdown.append('| '+' | '.join(row)+' |')
            markdown.append('')
        elif kind=='cover_art':
            # Small visual reference, not a decorative full-page cover.
            src=Path('/Users/seanpurtill/.codex/generated_images/01a08e7d-a646-7131-a24e-5276075de165/exec-e409bcd3-a9a0-4b76-97c3-d8fb079cf04f.png')
            if src.exists():
                pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;pp.add_run().add_picture(str(src),width=Inches(1.6))
        elif kind=='branch_diagram':
            pp=doc.add_paragraph();pp.add_run().add_picture(str(QA/'branch-diagram.png'),width=Inches(6.8))
            markdown.append('Diagram: local drying and smoking routes converge on preserved provisions; scholar teaching can supply smoking evidence. Containers, surplus, and carrying are also required for longer journeys.\n')
        elif kind=='atlas':
            pp=doc.add_paragraph();rr=pp.add_run(b[1]+'  ');rr.bold=True;rr=pp.add_run(b[2]);pp.paragraph_format.space_after=Pt(6)
            markdown.append(b[1]+'  '+b[2]+'\n')
        elif kind=='small':
            pp=doc.add_paragraph(b[1]);pp.paragraph_format.space_after=Pt(9)
            for rr in pp.runs:rr.font.size=Pt(9.5)
            markdown.append(b[1]+'\n')
        elif kind=='source':
            _,code,org,title,url,note=b
            pp=doc.add_paragraph();pp.add_run(f'[{code}] {org}. ').bold=True;link(pp,title,url=url)
            pp.add_run(' '+note);markdown.append(f'[{code}] {org}. [{title}]({url}). {note}\n')
        elif kind=='contents':
            for title_ref in b[1]:
                index=next(j for j,(tt,_) in enumerate(pages) if tt==title_ref)
                pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(10)
                link(pp,title_ref,anchor='section_'+str(index));pp.add_run('   '+str(index+1))
                markdown.append(title_ref+'\n')
    if title=='How to review this proposal':
        pp=doc.add_paragraph();link(pp,'Go to the candidate atlas',anchor='section_'+str(next(j for j,(tt,_) in enumerate(pages) if tt.startswith('Candidate atlas'))))
        pp.add_run('   |   ');link(pp,'Go to military forces',anchor='section_'+str(next(j for j,(tt,_) in enumerate(pages) if tt=='Military forces built from capabilities')))

# Remove built-in template paragraph borders, including the default blue Title rule.
for tree in (doc.styles.element,doc._element):
    for el in list(tree.iter(qn('w:pBdr'))):
        el.getparent().remove(el)
doc.save(OUT/'Tomorrow and Tomorrow Technology Design Review.docx')
(OUT/'Technology Design Review.md').write_text('\n'.join(markdown))
catalog=[]
for d in DOMAINS:
    for i,name in enumerate(d['entries'],1):
        catalog.append({'review_id':f"{d['id']}-{i:02}",'name':name,'field':d['name'],'status':'candidate for review','production_id':None})
(OUT/'candidate-atlas.json').write_text(json.dumps(catalog,indent=2)+'\n')
(QA/'page-plan.json').write_text(json.dumps([x[0] for x in pages],indent=2))
print(json.dumps({'planned_pages':len(pages),'candidate_discoveries':len(catalog),'target_discoveries':sum(d['target'] for d in DOMAINS),'military_roles':sum(len(x[1]) for x in MILITARY_ROLES),'words':len(' '.join(markdown).split()),'output':str(OUT/'Tomorrow and Tomorrow Technology Design Review.docx')},indent=2))
